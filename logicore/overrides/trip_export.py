import json
import frappe

_ADDRESS_SWAP = {
    "origin_address_1":      "origin_address_1_full",
    "origin_address_2":      "origin_address_2_full",
    "destination_address_1": "destination_address_1_full",
    "destination_address_2": "destination_address_2_full",
}


@frappe.whitelist()
def export_query():
    from frappe.desk.reportview import export_query as _orig

    # Only intercept Trip exports
    doctype = frappe.form_dict.get("doctype", "")
    if doctype != "Trip":
        return _orig()

    raw = frappe.form_dict.get("fields", "[]")
    try:
        requested = json.loads(raw) if isinstance(raw, str) else list(raw or [])
    except Exception:
        return _orig()

    swapped = {f: _ADDRESS_SWAP[f] for f in requested if f in _ADDRESS_SWAP}
    if not swapped:
        return _orig()

    # Silently replace address link fields with their stored full-text counterparts
    new_fields = [swapped.get(f, f) for f in requested]
    frappe.form_dict["fields"] = json.dumps(new_fields)

    # Run original export — builds file into frappe.response
    result = _orig()

    # Restore original column labels so user sees "Origin Address 1" not "Origin Address 1 Full"
    _fix_export_headers(swapped, frappe.form_dict.get("file_format_type", "CSV"))
    return result


def _fix_export_headers(swapped, file_format):
    content = frappe.response.get("filecontent")
    if not content:
        return

    meta = frappe.get_meta("Trip")
    label_swap = {}
    for orig_field, full_field in swapped.items():
        orig_f = meta.get_field(orig_field)
        full_f = meta.get_field(full_field)
        if orig_f and full_f:
            label_swap[full_f.label] = orig_f.label

    if not label_swap:
        return

    if file_format == "Excel":
        _fix_excel_headers(content, label_swap)
    else:
        _fix_csv_headers(content, label_swap)


def _fix_csv_headers(content, label_swap):
    try:
        text = content.decode("utf-8-sig") if isinstance(content, bytes) else content
        header, _, rest = text.partition("\n")
        for full_label, orig_label in label_swap.items():
            header = header.replace(f'"{full_label}"', f'"{orig_label}"')
        frappe.response["filecontent"] = (header + "\n" + rest).encode("utf-8-sig")
    except Exception:
        frappe.log_error(frappe.get_traceback(), "TMS: Failed to fix CSV export headers")


def _fix_excel_headers(content, label_swap):
    try:
        import io
        from openpyxl import load_workbook

        if isinstance(content, str):
            content = content.encode("utf-8")

        wb = load_workbook(io.BytesIO(content))
        ws = wb.active
        for col in range(1, (ws.max_column or 0) + 1):
            val = ws.cell(row=1, column=col).value
            if val in label_swap:
                ws.cell(row=1, column=col).value = label_swap[val]

        out = io.BytesIO()
        wb.save(out)
        frappe.response["filecontent"] = out.getvalue()
    except Exception:
        frappe.log_error(frappe.get_traceback(), "TMS: Failed to fix Excel export headers")


# ---------------------------------------------------------------------------
# Custom filtered Trip export (called from Trip list "Export Trips" button)
# ---------------------------------------------------------------------------

# Fields included in the export, in display order.
# address _full fields are used in place of the raw Link fields.
_TRIP_EXPORT_FIELDS = [
    ("name",                      "Trip No"),
    ("tcntrip_date",              "TCN/Trip Date"),
    ("company",                   "Company"),
    ("branch",                    "Branch"),
    ("business_format",           "Business Format"),
    ("trip_type",                 "Trip Type"),
    ("so_no",                     "Trip ID"),
    ("tcntrip_no",                "TCN/Trip No"),
    ("customer",                  "Customer"),
    ("vendor",                    "Vendor"),
    ("vehicle_no",                "Vehicle No"),
    ("vehicle_type",              "Vehicle Type"),
    ("driver",                    "Driver"),
    ("driver_mobile",             "Driver Mobile"),
    ("lr_no",                     "LR No"),
    ("lr_date",                   "LR Date"),
    ("origin_city",               "Origin City 1"),
    ("origin_city_2",             "Origin City 2"),
    ("destination_city_1",        "Destination City 1"),
    ("destination_city_2",        "Destination City 2"),
    ("origin_address_1_full",     "Origin Address 1"),
    ("origin_address_2_full",     "Origin Address 2"),
    ("destination_address_1_full","Destination Address 1"),
    ("destination_address_2_full","Destination Address 2"),
    ("billing_start_date",        "Billing Start Date"),
    ("billing_end_date",          "Billing End Date"),
    ("customer_freight",          "Customer Freight"),
    ("vendor_freight",            "Vendor Freight"),
    ("vendor_freight_extra",      "Vendor Freight (Extra)"),
    ("total_vendor_freight",      "Total Vendor Freight"),
    ("total_trip_amount",         "Trip Total Freight"),
    ("start_km",                  "Start KM"),
    ("end_km",                    "End KM"),
    ("total_km",                  "Total KM"),
    ("dispatch_date_time",        "Dispatch Date/Time"),
    ("pod_status",                "POD Status"),
    ("date_pod_uploaded_on_erp",  "Date POD Uploaded"),
    ("bill_nodate",               "Bill No"),
    ("bill_date",                 "Bill Date"),
    ("date_bill_credited",        "Date Bill Credited"),
    ("trip_status",               "Trip Status"),
    ("remarks",                   "Remarks"),
]


# Address Link → full-text field map (both directions for label lookup)
_ADDR_LINK_TO_FULL = {
    "origin_address_1":      "origin_address_1_full",
    "origin_address_2":      "origin_address_2_full",
    "destination_address_1": "destination_address_1_full",
    "destination_address_2": "destination_address_2_full",
}
_ADDR_FULL_TO_LINK = {v: k for k, v in _ADDR_LINK_TO_FULL.items()}

# Own-fleet field -> its free-text "Market" (non-fleet) counterpart. Both are
# selectable in the column picker, but the Market side never gets its own output
# column: whichever is selected (either, or both), the data always lands in the
# single fleet-named column (Vehicle No / Driver / Driver Mobile), falling back to
# the Market value when the fleet value is empty.
_MARKET_FALLBACK = {
    "vehicle_no":    "vehicle_market",
    "driver":        "driver_market",
    "driver_mobile": "driver_mobile_market",
}
_MARKET_TO_FLEET = {v: k for k, v in _MARKET_FALLBACK.items()}


def _collapse_market_columns(db_fields, headers):
    """Rewrite any Market-side field to its fleet counterpart and drop duplicates,
    so a pair is never exported as two columns."""
    meta = frappe.get_meta("Trip")
    out_fields, out_headers, seen = [], [], set()
    for fn, hd in zip(db_fields, headers):
        canon = _MARKET_TO_FLEET.get(fn, fn)
        if canon in seen:
            continue
        seen.add(canon)
        if canon != fn:
            df = meta.get_field(canon)
            hd = df.label if df else hd
        out_fields.append(canon)
        out_headers.append(hd)
    return out_fields, out_headers


def _fmt_val(v):
    from datetime import date, datetime
    if isinstance(v, datetime):
        return v.strftime("%d-%m-%Y %H:%M")
    if isinstance(v, date):
        return v.strftime("%d-%m-%Y")
    return v


def _resolve_link_fields(trip, fields_to_export):
    """Resolve Link field IDs to their display names."""
    # Resolve driver (Link → Driver) to driver's full_name
    if "driver" in fields_to_export and trip.get("driver"):
        driver_name = frappe.db.get_value("Driver", trip["driver"], "full_name")
        trip["driver"] = driver_name or trip["driver"]

    # Resolve trip_coordinator (Link → Employee) to employee's employee_name
    if "trip_coordinator" in fields_to_export and trip.get("trip_coordinator"):
        coordinator_name = frappe.db.get_value(
            "Employee", trip["trip_coordinator"], "employee_name"
        )
        trip["trip_coordinator"] = coordinator_name or trip["trip_coordinator"]

    return trip


def _build_frappe_filters(filters):
    """Convert our pre-built filter dict into a frappe.db filter list."""
    result = []

    # If specific record names were passed (list-view row selection), add as a base filter
    # Combined with any other form filters below (Python applies all conditions together)
    if filters.get("names"):
        names = filters["names"]
        if isinstance(names, list) and names:
            result.append(["Trip", "name", "in", names])

    def _add(field, op, val):
        if val:
            result.append(["Trip", field, op, val])

    _add("company",            "=",  filters.get("company"))
    _add("branch",             "=",  filters.get("branch"))
    _add("business_format",    "=",  filters.get("business_format"))
    _add("trip_type",          "=",  filters.get("trip_type"))
    _add("trip_status",        "=",  filters.get("trip_status"))
    _add("tcntrip_date",       ">=", filters.get("date_from"))
    _add("tcntrip_date",       "<=", filters.get("date_to"))
    _add("billing_start_date", ">=", filters.get("billing_start_date_from"))
    _add("billing_start_date", "<=", filters.get("billing_start_date_to"))
    _add("name",               ">=", filters.get("tcn_no_start"))
    _add("name",               "<=", filters.get("tcn_no_end"))
    _add("lr_no",              ">=", filters.get("lr_no_start"))
    _add("lr_no",              "<=", filters.get("lr_no_end"))
    _add("origin_city",        "=",  filters.get("origin_city"))
    _add("origin_city_2",      "=",  filters.get("origin_city_2"))
    _add("destination_city_1", "=",  filters.get("destination_city_1"))
    _add("destination_city_2", "=",  filters.get("destination_city_2"))
    _add("customer",           "=",  filters.get("customer"))
    _add("vendor",             "=",  filters.get("vendor"))
    _add("pod_status",         "=",  filters.get("pod_status"))
    _add("bill_date",          "=",  filters.get("bill_date"))
    _add("date_bill_credited", "=",  filters.get("date_bill_credited"))

    if filters.get("bill_no"):
        result.append(["Trip", "bill_nodate", "like", f"%{filters['bill_no']}%"])
    if filters.get("dispatch_date"):
        result.append(["Trip", "dispatch_date_time", "like", f"{filters['dispatch_date']}%"])

    return result


@frappe.whitelist()
def export_trips_filtered(filters=None, selected_fields=None):
    """
    Called via form POST from the Trip list Export Data dialog.
    - filters: dict of pre-built filter values from the injected filter section
    - selected_fields: {"Trip": [fieldname, ...]} from the MultiCheck column picker
    Address Link fields are automatically replaced with their _full counterparts.
    Always returns a CSV download.
    """
    import csv
    import io as _io
    import json

    if isinstance(filters, str):
        filters = json.loads(filters)
    filters = filters or {}

    if isinstance(selected_fields, str):
        selected_fields = json.loads(selected_fields)
    selected_fields = selected_fields or {}

    # ── Build filter conditions ──────────────────────────────────────────
    frappe_filters = _build_frappe_filters(filters)

    # ── Determine which columns to export ───────────────────────────────
    trip_fields_selected = selected_fields.get("Trip", [])

    if trip_fields_selected:
        # User chose specific columns from the MultiCheck.
        # Swap any address Link fields with their _full counterparts.
        db_fields = [_ADDR_LINK_TO_FULL.get(f, f) for f in trip_fields_selected]

        # Build headers: for _full fields use the original Link field's label
        meta = frappe.get_meta("Trip")

        def _label(fn):
            if fn == "name":
                return "ID"
            src = _ADDR_FULL_TO_LINK.get(fn, fn)   # resolve _full → original link
            df  = meta.get_field(src)
            return df.label if df else fn

        headers = [_label(fn) for fn in db_fields]
    else:
        # No columns chosen → export the default comprehensive set
        db_fields = [f[0] for f in _TRIP_EXPORT_FIELDS]
        headers   = [f[1] for f in _TRIP_EXPORT_FIELDS]

    # Vehicle No/Driver/Driver Mobile always export as one column, never split
    # into a separate "(Market)" column even if that checkbox was selected.
    db_fields, headers = _collapse_market_columns(db_fields, headers)

    # ── Query ────────────────────────────────────────────────────────────
    # Also fetch each Market fallback field alongside its own-fleet counterpart
    # (not added to db_fields/headers — used only to fill the same output column).
    fetch_fields = list(db_fields)
    for fn in db_fields:
        market_fn = _MARKET_FALLBACK.get(fn)
        if market_fn and market_fn not in fetch_fields:
            fetch_fields.append(market_fn)

    trips = frappe.db.get_all(
        "Trip",
        filters=frappe_filters or {},
        fields=fetch_fields,
        order_by="tcntrip_date desc, name desc",
        limit_page_length=0,
    )

    if not trips and frappe_filters:
        frappe.log_error(
            f"TMS Export returned 0 rows.\nFilters: {frappe_filters}\nFields: {db_fields}",
            "TMS Export: Zero Results"
        )

    # ── Build CSV ────────────────────────────────────────────────────────
    rows = [headers]
    for trip in trips:
        # Resolve Link field IDs to display names (driver, trip_coordinator)
        trip = _resolve_link_fields(trip, db_fields)

        def _cell(fn):
            val = trip.get(fn)
            if not val and fn in _MARKET_FALLBACK:
                val = trip.get(_MARKET_FALLBACK[fn])
            return val

        rows.append([
            str(_fmt_val(_cell(fn))) if _cell(fn) is not None else ""
            for fn in db_fields
        ])

    buf = _io.StringIO()
    csv.writer(buf).writerows(rows)

    ts = frappe.utils.now_datetime().strftime("%d-%m-%Y %H%M")
    frappe.response["filename"]    = f"Trip Export - {ts}.csv"
    frappe.response["filecontent"] = buf.getvalue().encode("utf-8-sig")
    frappe.response["type"]        = "binary"
